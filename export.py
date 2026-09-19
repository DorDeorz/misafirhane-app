# -*- coding: utf-8 -*-
"""
Excel'e aktarma işlemleri. openpyxl kullanır.
Tek yönlü aktarım: sistemdeki veriyi Excel'e yazar (rapor/yedek amaçlı).
Excel dosyası düzenlenip geri sisteme YÜKLENMEZ — asıl veri her zaman
misafirhane.db dosyasındadır.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime

import repository
from database import fiyat_tipi_goster


def _baslik_satiri_yaz(ws, basliklar):
    ws.append(basliklar)
    for col in range(1, len(basliklar) + 1):
        hucre = ws.cell(row=1, column=col)
        hucre.font = Font(bold=True, color="FFFFFF")
        hucre.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        hucre.alignment = Alignment(horizontal="center")


def rezervasyonlari_disa_aktar(dosya_yolu, durum="aktif", satirlar=None):
    """Tüm rezervasyon listesini (aktif/iptal/hepsi) Excel'e yazar.
    satirlar verilirse (örn. arama/sıralama uygulanmış özel bir liste),
    durum parametresi yerine doğrudan o liste kullanılır."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Rezervasyonlar"

    basliklar = [
        "ID", "Oda", "Ad Soyad", "TC No", "Telefon", "Kişi Sayısı",
        "Giriş Tarihi", "Toplam Gece", "Çıkış Tarihi", "Fiyat Tipi",
        "Toplam Tutar (TL)", "Referans", "Oda Sayısı", "Alan Kullanıcı",
        "Alınma Tarihi", "Durum", "Notlar"
    ]
    _baslik_satiri_yaz(ws, basliklar)

    rows = satirlar if satirlar is not None else repository.rezervasyon_listesi(durum)
    if satirlar is None and durum == "gelmedi":
        rows = [r for r in rows if (r["gelmedi_odasi"] or 0) > 0]
    ozet = repository.rezervasyonlari_toplam_ozeti([r["id"] for r in rows]) if rows else {}
    for r in rows:
        o = ozet.get(r["id"]) or {}
        toplam = o.get("toplam", 0)
        fiyat_birimleri = sorted({fiyat_tipi_goster(t) for t in o.get("fiyat_tipleri", set())})
        fiyat_metni = " + ".join(fiyat_birimleri) if fiyat_birimleri else "-"
        ws.append([
            r["id"], r["oda_ozeti"] or "-", r["ad_soyad"], r["tc_no"] or "",
            r["telefon"] or "", r["toplam_kisi"], r["giris_tarihi"], r["toplam_gece"],
            r["cikis_tarihi"] or "", fiyat_metni, toplam,
            r["referans"] or "", r["oda_sayisi"], r["olusturan_kullanici"] or "",
            r["olusturma_tarihi"] or "", r["durum_etiket"] or "", r["notlar"] or ""
        ])

    for col_letter, genislik in zip(
        "ABCDEFGHIJKLMNOPQ", [5, 28, 22, 14, 14, 8, 12, 10, 12, 12, 14, 18, 9, 14, 16, 16, 30]
    ):
        ws.column_dimensions[col_letter].width = genislik

    wb.save(dosya_yolu)
    return len(rows)


def gunluk_durumu_disa_aktar(dosya_yolu, tarih_str):
    """Belirli bir günün oda durumunu (doluluk + ödeme) Excel'e yazar."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"Oda Durumu {tarih_str}"

    basliklar = [
        "Kat", "Oda No", "Durum", "Ad Soyad", "TC No", "Telefon",
        "Kişi Sayısı", "Fiyat Tipi", "Gecelik Tutar", "Ödeme Durumu", "Ödeme Şekli"
    ]
    _baslik_satiri_yaz(ws, basliklar)

    rows = repository.gunun_oda_durumu(tarih_str)
    for r in rows:
        dolu_mu = r["rez_id"] is not None
        if dolu_mu:
            durum_metni = "Dolu"
        else:
            oda_durum = r.get("aktif_durum") or "temiz"
            if oda_durum == "temizlikte":
                durum_metni = "Boş (Temizlikte)"
            elif oda_durum == "arizali":
                durum_metni = "Boş (Arızalı)"
            else:
                durum_metni = "Boş"
        ws.append([
            r["kat_adi"], r["oda_no"], durum_metni,
            r["ad_soyad"] or "", r["tc_no"] or "", r["telefon"] or "",
            r["kisi_sayisi"] or "", fiyat_tipi_goster(r["fiyat_tipi"]) or "",
            r["tutar"] or "", "Ödendi" if r["odendi"] else ("Ödenmedi" if dolu_mu else ""),
            r["odeme_sekli"] or ""
        ])

    for col_letter, genislik in zip("ABCDEFGHIJK", [10, 8, 8, 22, 14, 14, 8, 10, 12, 12, 14]):
        ws.column_dimensions[col_letter].width = genislik

    wb.save(dosya_yolu)
    return len(rows)


def tarih_araligi_raporu_disa_aktar(dosya_yolu, baslangic_str, bitis_str):
    """Belirtilen tarih aralığındaki (baslangic ve bitis dahil) her gün için tüm
    odaların durumunu tek bir tabloda listeler (günlük/haftalık/aylık/özel rapor
    ihtiyaçlarının hepsi bu fonksiyonla karşılanır — tarih aralığı ne kadar geniş
    olursa o kadar kapsamlı rapor olur). Ayrıca özet bir sayfa da ekler."""
    from datetime import datetime as dt, timedelta as td

    baslangic = dt.strptime(baslangic_str, "%Y-%m-%d").date()
    bitis = dt.strptime(bitis_str, "%Y-%m-%d").date()
    if bitis < baslangic:
        baslangic, bitis = bitis, baslangic

    wb = Workbook()
    ws = wb.active
    ws.title = "Detay"

    basliklar = [
        "Tarih", "Kat", "Oda No", "Durum", "Ad Soyad", "TC No", "Telefon",
        "Kişi Sayısı", "Fiyat Tipi", "Gecelik Tutar", "Ödeme Durumu", "Ödeme Şekli"
    ]
    _baslik_satiri_yaz(ws, basliklar)

    toplam_dolu_gece = 0
    toplam_odenen = 0
    toplam_odenmeyen = 0
    toplam_gun_sayisi = (bitis - baslangic).days + 1

    gun = baslangic
    while gun <= bitis:
        gun_str = gun.isoformat()
        rows = repository.gunun_oda_durumu(gun_str)
        for r in rows:
            dolu_mu = r["rez_id"] is not None
            if dolu_mu:
                toplam_dolu_gece += 1
                if r["odendi"]:
                    toplam_odenen += r["tutar"] or 0
                else:
                    toplam_odenmeyen += r["tutar"] or 0
            if dolu_mu:
                durum_metni = "Dolu"
            else:
                oda_durum = r.get("aktif_durum") or "temiz"
                if oda_durum == "temizlikte":
                    durum_metni = "Boş (Temizlikte)"
                elif oda_durum == "arizali":
                    durum_metni = "Boş (Arızalı)"
                else:
                    durum_metni = "Boş"
            ws.append([
                gun_str, r["kat_adi"], r["oda_no"], durum_metni,
                r["ad_soyad"] or "", r["tc_no"] or "", r["telefon"] or "",
                r["kisi_sayisi"] or "", fiyat_tipi_goster(r["fiyat_tipi"]) or "",
                r["tutar"] or "", "Ödendi" if r["odendi"] else ("Ödenmedi" if dolu_mu else ""),
                r["odeme_sekli"] or ""
            ])
        gun += td(days=1)

    for col_letter, genislik in zip("ABCDEFGHIJKL", [12, 10, 8, 8, 22, 14, 14, 8, 10, 12, 12, 14]):
        ws.column_dimensions[col_letter].width = genislik

    # Ozet sayfasi
    ws2 = wb.create_sheet("Özet")
    _baslik_satiri_yaz(ws2, ["Bilgi", "Değer"])
    ws2.append(["Başlangıç Tarihi", baslangic_str])
    ws2.append(["Bitiş Tarihi", bitis_str])
    ws2.append(["Toplam Gün Sayısı", toplam_gun_sayisi])
    ws2.append(["Toplam Dolu Oda-Gece", toplam_dolu_gece])
    ws2.append(["Tahsil Edilen Toplam Tutar (TL)", toplam_odenen])
    ws2.append(["Tahsil Edilmemiş Toplam Tutar (TL)", toplam_odenmeyen])
    ws2.append(["Beklenen Toplam Gelir (TL)", toplam_odenen + toplam_odenmeyen])
    for col_letter, genislik in zip("AB", [30, 20]):
        ws2.column_dimensions[col_letter].width = genislik

    wb.save(dosya_yolu)
    return toplam_gun_sayisi
