# -*- coding: utf-8 -*-
"""
1.0.6 (Konaklayan Listesi) için izole test: iş katmanı + offscreen arayüz.

Kapsam: eski notların not geçmişine taşınması, not ekleme/silme (yazan ve
zamanıyla), misafir kartı (puan, sorunlu nedeni ve işaretleyen), geldiği yer,
misafir listesi (kişi birleştirme, arama), referans listesi ve notları,
konaklayan listesi (tarih aralığı, oda değiştiren misafir tek satır),
Excel (puan/sorunlu yalnızca istenince), yazdırma HTML'i, Misafirler sekmesi
ve rezervasyon detayındaki not kutuları.

Kendi GEÇİCİ veritabanını kurar (TEMP altında); gerçek veriye DOKUNMAZ.
Çalıştırma: python konaklayan_listesi_test.py
"""
import os
import sys
import shutil
import sqlite3
import tempfile
from datetime import date, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

TEST_KLASOR = os.path.join(tempfile.gettempdir(), "opencode", "konaklayan_listesi_test")
if os.path.isdir(TEST_KLASOR):
    shutil.rmtree(TEST_KLASOR)
os.makedirs(TEST_KLASOR, exist_ok=True)

import database

database.VERI_KLASORU = TEST_KLASOR
database.DB_PATH = os.path.join(TEST_KLASOR, "misafirhane.db")

# 1.0.5 şemasının ilgili parçası: tek metinlik rezervasyon notu ve kart notu.
_c = sqlite3.connect(database.DB_PATH)
_c.execute("CREATE TABLE rezervasyonlar (id INTEGER PRIMARY KEY AUTOINCREMENT, ad_soyad TEXT NOT NULL, "
           "tc_no TEXT, telefon TEXT, referans TEXT, notlar TEXT DEFAULT '', olusturan_kullanici TEXT, "
           "iptal INTEGER DEFAULT 0, iptal_nedeni TEXT, olusturma_tarihi TEXT DEFAULT CURRENT_TIMESTAMP)")
_c.execute("INSERT INTO rezervasyonlar (ad_soyad, telefon, notlar, olusturan_kullanici) "
           "VALUES ('Eski Kayit', '0533 999 88 77', 'eski rezervasyon notu', 'ayse')")
_c.execute("CREATE TABLE misafir_kartlari (id INTEGER PRIMARY KEY AUTOINCREMENT, telefon_anahtar TEXT DEFAULT '', "
           "tc_no TEXT DEFAULT '', ad_soyad TEXT DEFAULT '', notu TEXT DEFAULT '', sorunlu INTEGER DEFAULT 0, "
           "guncelleyen TEXT, guncelleme_zamani TEXT)")
_c.execute("INSERT INTO misafir_kartlari (telefon_anahtar, ad_soyad, notu, sorunlu, guncelleyen, guncelleme_zamani) "
           "VALUES ('5339998877', 'Eski Kayit', 'borcunu ödemedi', 1, 'mehmet', '2026-01-05 10:00:00')")
_c.commit()
_c.close()

database.init_db()
database.init_db()  # ikinci açılış notları tekrar taşımamalı
database.varsayilan_odalari_yukle()

import repository as R
import loglama

B = date.today()
hatalar = []


def g(n):
    return (B + timedelta(days=n)).isoformat()


def kontrol(kosul, mesaj):
    if not kosul:
        hatalar.append(mesaj)
        print("  [HATA]", mesaj)
    else:
        print("  [ok]  ", mesaj)


def hata_verir_mi(f):
    try:
        f()
    except ValueError:
        return True
    return False


O = {o["oda_no"]: o["id"] for o in R.oda_listesi()}
YER = dict(tc_no="", fiyat_tipi="Sabit", gecelik_ucret=1300)


def rez(oda_no, ofs, gece, ad, telefon="", referans="", yer="", notlar="", kullanici="ayse", kisi=1):
    rid = R.rezervasyon_olustur([dict(oda_id=O[oda_no], giris_tarihi=g(ofs), gece_sayisi=gece,
                                      kisi_sayisi=kisi, fiyat_tipi="Sabit")], ad, telefon=telefon,
                                referans=referans, notlar=notlar, olusturan_kullanici=kullanici,
                                gecmis_kontrol=False, geldigi_yer=yer)
    return rid, R.rezervasyon_odalar_listele(rid)[0]["id"]


def checkin(ro_id, *kisiler):
    conn = database.get_connection()
    conn.execute("UPDATE rezervasyon_odalar SET checkin_yapildi=1 WHERE id=?", (ro_id,))
    conn.commit()
    conn.close()
    R.odasi_misafirleri_kaydet(ro_id, [dict(YER, ad_soyad=a, tc_no=t) for a, t in kisiler])


print("Eski notların taşınması")
eski = R.notlar_listele("rezervasyon", 1)
kontrol(len(eski) == 1 and eski[0]["metin"] == "eski rezervasyon notu" and eski[0]["yazan"] == "ayse",
        "rezervasyon notu, alan kullanıcıyla not geçmişine taşındı (bir kez)")
eski_kart = R.misafir_karti_getir(telefon="0533 999 88 77")
kontrol(eski_kart["sorunlu_nedeni"] == "borcunu ödemedi" and eski_kart["sorunlu_isaretleyen"] == "mehmet",
        "sorunlu kartın notu, sorunlu nedeni oldu")
kontrol([n["metin"] for n in R.notlar_listele("misafir", eski_kart["id"])] == ["borcunu ödemedi"],
        "kart notu misafir notu oldu")

print("Notlar")
loglama.set_aktif_kullanici("mehmet")
r1, ro1 = rez(1, -10, 3, "Ali Veli", "0532 111 22 33", "Başkan Ahmet Bey", "Ankara", "geç gelecek")
n = R.notlar_listele("rezervasyon", r1)
kontrol(len(n) == 1 and n[0]["yazan"] == "ayse", "rezervasyon alınırken yazılan not, alanın adıyla kaydedilir")
R.not_ekle("rezervasyon", r1, "oda değişikliği istedi")
n = R.notlar_listele("rezervasyon", r1)
kontrol(n[0]["metin"] == "oda değişikliği istedi" and n[0]["yazan"] == "mehmet" and n[0]["zaman"],
        "yeni not giriş yapan kullanıcı ve zamanla eklenir, en yeni üstte")
kontrol(hata_verir_mi(lambda: R.not_ekle("rezervasyon", r1, "   ")), "boş not reddedilir")
kontrol(hata_verir_mi(lambda: R.not_ekle("bilinmeyen", r1, "x")), "bilinmeyen not türü reddedilir")
R.not_sil(n[0]["id"])
kontrol(len(R.notlar_listele("rezervasyon", r1)) == 1, "not silinir")
kontrol(any(i["tur"] == "not_sil" for i in loglama.son_islemler(20)), "silinen not işlem geçmişine yazılır")
R.not_ekle("referans", "başkan  ahmet BEY", "hafta sonu gelirler")
kontrol(len(R.notlar_listele("referans", "Başkan Ahmet Bey")) == 1, "referans notu harf/boşluk farkından bağımsız")
kontrol(R.rezervasyon_getir(r1)["geldigi_yer"] == "Ankara", "geldiği yer kaydedilir")
R.rezervasyon_guncelle(r1, "Ali Veli", "", "0532 111 22 33", "Başkan Ahmet Bey", geldigi_yer="Konya")
kontrol(R.rezervasyon_getir(r1)["geldigi_yer"] == "Konya", "geldiği yer güncellenir")
R.rezervasyon_guncelle(r1, "Ali Veli", "", "0532 111 22 33", "Başkan Ahmet Bey", geldigi_yer="Ankara")

print("Misafir kartı: puan ve sorunlu nedeni")
kontrol(hata_verir_mi(lambda: R.misafir_karti_kaydet("0532 111 22 33", "", "Ali Veli", 6)), "puan 1-5 arası")
kart = R.misafir_karti_olustur("0532 111 22 33", "", "Ali Veli")
R.not_ekle("misafir", kart["id"], "çay sever")
R.misafir_karti_kaydet("0532 111 22 33", "", "Ali Veli", 4, True, "gürültü yaptı")
kart = R.misafir_karti_getir(telefon="05321112233")
kontrol(kart["puan"] == 4 and kart["sorunlu"] and kart["sorunlu_nedeni"] == "gürültü yaptı"
        and kart["sorunlu_isaretleyen"] == "mehmet" and kart["sorunlu_zamani"], "işaretleyen ve zaman yazılır")
loglama.set_aktif_kullanici("ayse")
R.misafir_karti_kaydet("0532 111 22 33", "", "Ali Veli", 3, True, "gürültü yaptı")
kontrol(R.misafir_karti_getir(telefon="05321112233")["sorunlu_isaretleyen"] == "mehmet",
        "yalnız puan değişince işaretleyen değişmez")
R.misafir_karti_kaydet("0532 111 22 33", "", "Ali Veli", 0, False)
kart = R.misafir_karti_getir(telefon="05321112233")
kontrol(kart is not None and not kart["sorunlu"] and kart["sorunlu_nedeni"] == "",
        "notu olan kart puansız/sorunsuz kalınca silinmez")
R.misafir_karti_kaydet("0532 111 22 33", "", "Ali Veli", 4, True, "gürültü yaptı")  # ayse yeniden işaretledi
loglama.set_aktif_kullanici("mehmet")

print("Misafir listesi")
checkin(ro1, ("Ali Veli", "10000000146"))
r2, ro2 = rez(2, -2, 4, "ali  veli", "+90 532 111 2233", "başkan ahmet  bey", "")
checkin(ro2, ("Ali Veli", "10000000146"))
r3, ro3 = rez(3, 5, 2, "Can Ileri", "0544 000 00 00", "Başkan Ahmet Bey")
r4, ro4 = rez(4, -5, 2, "Zeynep Öz", "0555 123 45 67", "", "İzmir", kisi=2)
checkin(ro4, ("Zeynep Öz", "10000000528"), ("Hasan Öz", "10000000900"))
liste = R.misafir_listesi()
kontrol(len(liste) == 2, "yalnız konaklamış kişiler, aynı telefon tek kişi")
ali = [m for m in liste if m["telefon"].startswith("+90")][0]
kontrol(ali["konaklama"] == 2 and ali["toplam_gece"] == 7 and ali["ad_soyad"] == "ali veli",
        "konaklama ve gece sayısı, son yazılan ad")
kontrol(ali["geldigi_yer"] == "Ankara" and ali["puan"] == 4 and ali["sorunlu"] and ali["not_sayisi"] == 1,
        "geldiği yer, puan, sorunlu ve not sayısı")
kontrol(liste[0]["son_giris"] >= liste[1]["son_giris"], "son konaklamaya göre sıralı")
kontrol([m["ad_soyad"] for m in R.misafir_listesi("zeynep")] == ["Zeynep Öz"], "adla arama")
kontrol(len(R.misafir_listesi("İZMİR")) == 1, "geldiği yerle arama, Türkçe büyük harf")
kontrol(len(R.misafir_listesi("532 111")) == 1 and not R.misafir_listesi("12"), "telefonla arama")
kontrol(len(R.konaklama_ozetleri(ali["rez_idler"])) == 2, "misafirin konaklama özetleri")
kontrol("Ankara" in R.gecmis_geldigi_yerler(), "geldiği yer otomatik tamamlama listesi")

print("Referanslar")
ref = R.referans_listesi()
kontrol(len(ref) == 1 and ref[0]["rezervasyon"] == 3 and ref[0]["konaklayan"] == 2 and ref[0]["not_sayisi"] == 1,
        "aynı referans tek satır, sayılar ve not")
kontrol(len(R.referans_rezervasyonlari("BAŞKAN AHMET BEY")) == 3, "referansla gelenler (gelmemiş dahil)")
R.rezervasyon_iptal(r3)
kontrol(R.referans_listesi()[0]["rezervasyon"] == 2, "iptal edilen sayılmaz")

print("Konaklayan listesi")
k = R.konaklayan_listesi(g(-30), g(0))
kontrol(len(k) == 4, "aralıkta kalan her kişi ayrı satır")
kontrol([x["giris"] for x in k] == sorted(x["giris"] for x in k), "girişe göre sıralı")
z = [x for x in k if x["ad_soyad"] == "Zeynep Öz"][0]
kontrol(z["telefon"] == "0555 123 45 67" and z["geldigi_yer"] == "İzmir" and z["gece"] == 2 and z["alan"] == "ayse",
        "telefon, geldiği yer, gece ve kaydı alan")
kontrol(len(R.konaklayan_listesi(g(-1), g(-1))) == 1, "tek gün: yalnız o gece kalan")
kontrol(len(R.konaklayan_listesi(g(-4), g(-3))) == 2, "aralığa kısmen giren konaklama dahil")
kontrol(not R.konaklayan_listesi(g(10), g(20)), "gelecekte (check-in yok) boş")
R.oda_degistir(ro2, O[5], g(0))
k = [x for x in R.konaklayan_listesi(g(-30), g(5)) if x["rez_id"] == r2]
kontrol(len(k) == 1 and k[0]["odalar"] == "2 → 5" and k[0]["giris"] == g(-2) and k[0]["cikis"] == g(2),
        "oda değiştiren misafir tek satır (2 → 5)")
kontrol(k[0]["puan"] == 4 and k[0]["sorunlu"], "puan/sorunlu kişiye bağlanır")
kontrol(hata_verir_mi(lambda: R.konaklayan_listesi("bozuk", g(0))), "bozuk tarih reddedilir")

print("Excel")
import export
from openpyxl import load_workbook
dosya = os.path.join(TEST_KLASOR, "liste.xlsx")
kontrol(export.konaklayan_listesi_disa_aktar(dosya, g(-30), g(5)) == 4, "Excel kişi sayısı")
ws = load_workbook(dosya).active
basliklar = [c.value for c in ws[3]]
kontrol("Puan" not in basliklar and "Sorunlu" not in basliklar, "varsayılan: puan/sorunlu yok")
kontrol(ws.freeze_panes == "A4" and ws.page_setup.orientation == "landscape", "başlık sabit, yatay sayfa")
tum = " ".join(str(c.value) for r in ws.iter_rows() for c in r if c.value)
kontrol("gürültü" not in tum, "sorunlu nedeni dosyada yok")
export.konaklayan_listesi_disa_aktar(dosya, g(-30), g(5), hassas=True)
ws = load_workbook(dosya).active
kontrol([c.value for c in ws[3]][-2:] == ["Puan", "Sorunlu"], "istenince puan/sorunlu sütunları eklenir")
kontrol(any("gürültü" in str(c.value) for r in ws.iter_rows() for c in r if c.value), "sorunlu nedeni yazılır")
R.not_ekle("rezervasyon", r4, "=HYPERLINK(\"x\")")
export.rezervasyonlari_disa_aktar(dosya, "hepsi")
ws = load_workbook(dosya).active
kontrol(any(str(c.value).startswith("'=") for r in ws.iter_rows() for c in r if c.value),
        "rezervasyon Excel'inde notlar (formül koruması ile)")

print("Arayüz (offscreen)")
from PySide6.QtWidgets import QApplication, QMessageBox
uyg = QApplication.instance() or QApplication(sys.argv)
import misafir_pencere
html = misafir_pencere.konaklayan_listesi_html(g(-30), g(5), R.konaklayan_listesi(g(-30), g(5)))
kontrol("Zeynep Öz" in html and "gürültü" not in html, "yazdırma HTML'i (hassas bilgi yok)")
kontrol("gürültü" in misafir_pencere.konaklayan_listesi_html(
    g(-30), g(5), R.konaklayan_listesi(g(-30), g(5)), True), "yazdırmada istenince hassas bilgi")

t = misafir_pencere.MisafirlerTab()
kontrol(t.misafir_tablo.rowCount() == 2, "Misafirler sekmesi dolar")
t.arama.setText("zeynep")
t._misafirleri_doldur()
kontrol(t.misafir_tablo.rowCount() == 1, "sekmede arama")
t.arama.setText("")
t.yalniz_sorunlu.setChecked(True)
kontrol(t.misafir_tablo.rowCount() == 1, "yalnız sorunlular filtresi")
t.yalniz_sorunlu.setChecked(False)
t.sekmeler.setCurrentIndex(1)
t._bu_ay()
kontrol(t.k_tablo.rowCount() >= 1 and t.k_tablo.isColumnHidden(12), "konaklayan listesi, puan sütunu gizli")
t.hassas.setChecked(True)
kontrol(not t.k_tablo.isColumnHidden(12), "kutucuk işaretlenince puan sütunu görünür")
t._bugun()
kontrol(t.k_bas.date() == t.k_bit.date(), "Bugün düğmesi")
t.sekmeler.setCurrentIndex(2)
kontrol(t.ref_tablo.rowCount() == 1, "referanslar alt sekmesi")

md = misafir_pencere.MisafirDetayDialog(ali)
kontrol(md.konaklamalar.rowCount() == 2 and md.kart_kutusu.puan.currentData() == 4
        and md.kart_kutusu.sorunlu.isChecked() and "ayse" in md.kart_kutusu.isaretleyen.text(),
        "misafir kartı penceresi: konaklamalar, puan, sorunlu ve işaretleyen")
kontrol(md.notlar.not_sayisi() == 1, "misafir kartında notlar")
md.notlar.yeni.setText("yine geldi")
md.notlar.ekle()
kontrol(md.notlar.not_sayisi() == 2 and md.degisti, "pencereden not eklenir")
md.kart_kutusu.puan.setCurrentIndex(5)
md.kaydet()
kontrol(R.misafir_karti_getir(telefon="05321112233")["puan"] == 5, "pencereden puan kaydedilir")
_orj_w = QMessageBox.warning
QMessageBox.warning = staticmethod(lambda *a, **k: None)
md = misafir_pencere.MisafirDetayDialog(R.misafir_listesi("zeynep")[0])
md.kart_kutusu.sorunlu.setChecked(True)
md.kaydet()
kontrol(not R.misafir_karti_getir(telefon="0555 123 45 67"), "nedensiz sorunlu işareti kaydedilmez")
md.notlar.yeni.setText("ilk not")
md.notlar.ekle()
kontrol(R.misafir_karti_getir(telefon="0555 123 45 67") is not None, "kartı olmayan misafire not eklenince kart açılır")
QMessageBox.warning = _orj_w
rd = misafir_pencere.ReferansDialog("Başkan Ahmet Bey")
kontrol(rd.tablo.rowCount() == 2 and rd.notlar.not_sayisi() == 1, "referans penceresi")

from detay_dialog import RezervasyonDetayDialog
dd = RezervasyonDetayDialog(r1)
kontrol(dd.rez_notlari.not_sayisi() == 1 and dd.misafir_notlari.not_sayisi() == 2
        and dd.referans_notlari.not_sayisi() == 1, "detayda üç not kutusu")
kontrol(dd.geldigi_yer.text() == "Ankara" and dd.kart_kutusu.puan.currentData() == 5, "detayda geldiği yer ve puan")
dd.rez_notlari.yeni.setText("detaydan not")
dd.rez_notlari.ekle()
kontrol(R.notlar_listele("rezervasyon", r1)[0]["metin"] == "detaydan not", "detaydan not hemen kaydedilir")
dd.geldigi_yer.setText("Sivas")
dd.kaydet()
kontrol(R.rezervasyon_getir(r1)["geldigi_yer"] == "Sivas", "detaydan geldiği yer kaydedilir")
r5, _ = rez(6, 3, 1, "Telefonsuz")
dd = RezervasyonDetayDialog(r5)
kontrol(not dd.misafir_notlari.yeni.isEnabled() and not dd.referans_notlari.yeni.isEnabled(),
        "telefonsuz/referanssız kayıtta misafir ve referans notu kapalı")

import main
yr = main.YeniRezervasyonTab()
yr.telefon.setText("0532 111 22 33")
yr.referans.setText("Başkan Ahmet Bey")
yr._misafiri_tani()
metin = yr.misafir_bilgi.text()
kontrol(yr._misafir_sorunlu and "gürültü yaptı" in yr._misafir_sorun_metni and "ayse" in yr._misafir_sorun_metni,
        "yeni rezervasyonda sorunlu nedeni ve işaretleyen")
kontrol("yine geldi" in metin and "hafta sonu gelirler" in metin and "★★★★★" in metin,
        "yeni rezervasyonda misafir notları, referans notu ve puan")
kontrol(yr.geldigi_yer.text() == "Sivas", "geldiği yer son konaklamadan dolar")

print("Sol kenar çubuğu")
main.AnaPencere._acilis_akisi = lambda self: None
ana = main.AnaPencere()
k = ana.kenar
kontrol(len(k._sayfa_butonlari) == ana.tabs.count() == 12 and ana.tabs.tabBar().isHidden(),
        "her sekme için kenar düğmesi, üstteki sekme çubuğu gizli")
k._sayfa_butonlari[6].click()
kontrol(ana.tabs.currentWidget() is ana.misafirler_tab, "kenar düğmesi sekmeyi açar")
ana.tabs.setCurrentIndex(2)
kontrol(k._sayfa_butonlari[2].isChecked() and not k._sayfa_butonlari[6].isChecked(), "seçili düğme sekmeyle eşleşir")
kontrol(len(k._tum_butonlar) == 12 + 5, "araç pencereleri de kenarda (Takvim, Excel, KBS, Kasa, Özet)")
k.daralt(True, animasyonlu=False)
kontrol(k.width() <= k.DAR and k._sayfa_butonlari[0].text() == "➕"
        and "Yeni Rezervasyon" in k._sayfa_butonlari[0].toolTip(), "daralınca yalnız simge, ad ipucunda")
kontrol(database.get_ayar("kenar_cubugu_dar") == "1", "dar durum kaydedilir")
kontrol(main.AnaPencere().kenar.dar, "uygulama aynı (dar) halde açılır")
k.daralt(False, animasyonlu=False)
kontrol(k.maximumWidth() == k.GENIS and "Yeni Rezervasyon" in k._sayfa_butonlari[0].text(), "genişletilir")

print()
if hatalar:
    print(f"{len(hatalar)} HATA:")
    for h in hatalar:
        print(" -", h)
    sys.exit(1)
print("TUM KONAKLAYAN LISTESI TESTLERI GECTI")
