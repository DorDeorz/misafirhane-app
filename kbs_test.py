# -*- coding: utf-8 -*-
"""
kbs.py için izole CLI testi.

Küçük bir GEÇİCİ veritabanı üretir (temp/opencode/kbs_test/), yerli + yabancı
+ çıkış yapılmış durumları kurar, modülü çalıştırır ve Excel çıktısını doğrular.
Uygulamanın gerçek verisine ve koduna DOKUNMAZ.
"""

import os
import sqlite3
import tempfile

import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import kbs

TEST_KLASOR = os.path.join(
    os.environ.get("TEMP", tempfile.gettempdir()), "opencode", "kbs_test")
os.makedirs(TEST_KLASOR, exist_ok=True)
DB = os.path.join(TEST_KLASOR, "test.db")
TAKIP = os.path.join(TEST_KLASOR, "kbs_takip.db")
EXCEL = os.path.join(TEST_KLASOR, "KBS_Bildirim_Test.xlsx")
for f in (DB, TAKIP, EXCEL):
    if os.path.exists(f):
        os.remove(f)


def tc_uret(taban="100000001"):
    """Geçerli bir T.C. kimlik no üretir (sağlama algoritmasıyla)."""
    haneler = [int(h) for h in taban]
    tek = sum(haneler[0::2][:5])
    cift = sum(haneler[1::2])
    d10 = (tek * 7 - cift) % 10
    d11 = (sum(haneler) + d10) % 10
    return taban + str(d10) + str(d11)


def kurulum():
    conn = sqlite3.connect(DB)
    cur = conn.cursor()
    cur.executescript("""
        CREATE TABLE odalar (
            id INTEGER PRIMARY KEY AUTOINCREMENT, eski_no INTEGER, kat_no INTEGER,
            kat_adi TEXT, oda_no INTEGER, oda_tipi TEXT, kapasite INTEGER DEFAULT 1,
            aktif INTEGER DEFAULT 1, durum TEXT DEFAULT 'temiz', ariza_bitis TEXT);
        CREATE TABLE rezervasyonlar (
            id INTEGER PRIMARY KEY AUTOINCREMENT, ad_soyad TEXT NOT NULL, tc_no TEXT,
            telefon TEXT, referans TEXT, notlar TEXT DEFAULT '',
            olusturan_kullanici TEXT, iptal INTEGER DEFAULT 0,
            olusturma_tarihi TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE rezervasyon_odalar (
            id INTEGER PRIMARY KEY AUTOINCREMENT, rezervasyon_id INTEGER NOT NULL,
            oda_id INTEGER NOT NULL, giris_tarihi TEXT NOT NULL,
            gece_sayisi INTEGER NOT NULL DEFAULT 1, cikis_tarihi TEXT,
            kisi_sayisi INTEGER NOT NULL DEFAULT 1, fiyat_tipi TEXT DEFAULT 'Sabit',
            gecelik_ucret INTEGER DEFAULT 1300, checkin_yapildi INTEGER DEFAULT 0);
        CREATE TABLE misafirler (
            id INTEGER PRIMARY KEY AUTOINCREMENT, rezervasyon_oda_id INTEGER NOT NULL,
            ad_soyad TEXT, tc_no TEXT, sira_no INTEGER DEFAULT 1,
            fiyat_tipi TEXT, gecelik_ucret INTEGER,
            uyruk TEXT DEFAULT '', dogum_tarihi TEXT DEFAULT '',
            cinsiyet TEXT DEFAULT '', dogum_yeri TEXT DEFAULT '',
            belge_turu TEXT DEFAULT '');
    """)
    cur.execute("INSERT INTO odalar (kat_adi, oda_no, oda_tipi) VALUES ('Lobi', 1, 'Tek')")
    cur.execute("INSERT INTO odalar (kat_adi, oda_no, oda_tipi) VALUES ('Lobi', 2, 'Double')")
    cur.execute("INSERT INTO odalar (kat_adi, oda_no, oda_tipi) VALUES ('1.KAT', 3, 'Aile')")

    # Rez 1: yerli, check-in yapılmış, henüz çıkış yok
    cur.execute("INSERT INTO rezervasyonlar (ad_soyad, telefon) VALUES ('Ahmet Yılmaz', '05551112233')")
    cur.execute("INSERT INTO rezervasyon_odalar (rezervasyon_id, oda_id, giris_tarihi, gece_sayisi, checkin_yapildi)"
                " VALUES (1, 1, '2026-09-18', 3, 1)")
    cur.execute("INSERT INTO misafirler (rezervasyon_oda_id, ad_soyad, tc_no, sira_no) "
                "VALUES (1, 'Ahmet Yılmaz', ?, 1)", (tc_uret(),))

    # Rez 2: yerli, BİLEREK bozuk TC (sağlama hatası) + çıkış yapılmış
    cur.execute("INSERT INTO rezervasyonlar (ad_soyad) VALUES ('Ayşe Demir')")
    cur.execute("INSERT INTO rezervasyon_odalar (rezervasyon_id, oda_id, giris_tarihi, gece_sayisi, checkin_yapildi, cikis_tarihi)"
                " VALUES (2, 2, '2026-09-15', 3, 1, '2026-09-18')")
    cur.execute("INSERT INTO misafirler (rezervasyon_oda_id, ad_soyad, tc_no, sira_no) "
                "VALUES (2, 'Ayşe Demir', '11111111111', 1)")

    # Rez 3: yabancı (belge no 'yabancı' sayılır), check-in yapılmış, bilgisi TAM
    cur.execute("INSERT INTO rezervasyonlar (ad_soyad, telefon) VALUES ('John Smith', '+905001112233')")
    cur.execute("INSERT INTO rezervasyon_odalar (rezervasyon_id, oda_id, giris_tarihi, gece_sayisi, checkin_yapildi)"
                " VALUES (3, 3, '2026-09-19', 2, 1)")
    cur.execute("INSERT INTO misafirler (rezervasyon_oda_id, ad_soyad, tc_no, sira_no,"
                " uyruk, dogum_tarihi, cinsiyet, dogum_yeri, belge_turu) "
                "VALUES (3, 'John Smith', 'A12345678', 1, 'Alman', '1985-04-12', 'Erkek', 'Hamburg', 'Pasaport')")

    # Rez 4: yabancı, bilgileri EKSİK (uyruk var, diğerleri yok) → uyarı göstermeli
    cur.execute("INSERT INTO rezervasyonlar (ad_soyad) VALUES ('Petra Schmidt')")
    cur.execute("INSERT INTO rezervasyon_odalar (rezervasyon_id, oda_id, giris_tarihi, gece_sayisi, checkin_yapildi)"
                " VALUES (4, 1, '2026-09-19', 2, 1)")
    cur.execute("INSERT INTO misafirler (rezervasyon_oda_id, ad_soyad, tc_no, sira_no, uyruk) "
                "VALUES (4, 'Petra Schmidt', 'C54321678', 1, 'Rus')")

    conn.commit()
    conn.close()


def main():
    kurulum()
    hatalar = []

    # 1) Yerli/yabancı sınıflandırma ve TC sağlama
    assert kbs.tanitim_kodu_gecerli_mi(tc_uret()) == "yerli"
    assert kbs.tanitim_kodu_gecerli_mi("A12345678") == "yabanci"
    assert kbs.tanitim_kodu_gecerli_mi("") == "eksik"
    assert kbs.tc_dogrula(tc_uret()) is True
    assert kbs.tc_dogrula("11111111111") is False

    # 2) Bekleyen listeleri
    giris, cikis = kbs.kbs_bekleyenler(DB, TAKIP)
    print("GİRİŞ bekleyen:", len(giris), "| ÇIKIŞ bekleyen:", len(cikis))
    if len(giris) != 4:
        hatalar.append("giriş bekleyen beklenen 4, bulunan %d" % len(giris))
    if len(cikis) != 1:
        hatalar.append("çıkış bekleyen beklenen 1, bulunan %d" % len(cikis))
    for s in giris:
        print("  [GİRİŞ]", s["tarih"], s["tc_no"], s["misafir_ad"], "|", s["oda"], "|", s["not"])
    for s in cikis:
        print("  [ÇIKIŞ]", s["tarih"], s["tc_no"], s["misafir_ad"], "|", s["oda"], "|", s["not"])

    john = [s for s in giris if s["misafir_ad"] == "John Smith"]
    if len(john) != 1 or john[0]["tip"] != "yabanci":
        hatalar.append("John Smith yabancı tespit edilemedi")
    elif john[0]["uyruk"] != "Alman" or john[0]["dogum_yeri"] != "Hamburg":
        hatalar.append("John Smith bilgi tam alanları taşınamadı (" + str(john[0]["uyruk"]) + ")")
    elif "tam bilgi zorunlu" in john[0]["not"]:
        hatalar.append("bilgisi tam yabancı yine de 'tam bilgi zorunlu' uyarısı aldı")

    petra = [s for s in giris if s["misafir_ad"] == "Petra Schmidt"]
    if len(petra) != 1 or "DOĞUM TARİHİ" not in petra[0]["not"]:
        hatalar.append("eksik bilgili yabancı gerekli uyarıyı almadı")

    bozuk_tc = [s for s in giris if "SAĞLAMA HATASI" in s.get("not", "")]
    if not bozuk_tc:
        hatalar.append("bozuk TC işaretlenmedi")

    # 3) Excel çıktısı
    ig, ic = kbs.kbs_excel_ure(EXCEL, DB, TAKIP)
    if not os.path.exists(EXCEL):
        hatalar.append("Excel üretilemedi")
    else:
        from openpyxl import load_workbook
        wb = load_workbook(EXCEL)
        sayfalar = wb.sheetnames
        ws = wb["GİRİŞ BİLDİRİMLERİ"]
        veri_satiri = ws.max_row - 1
        print("Excel:", os.path.basename(EXCEL), "| sayfalar:", sayfalar,
              "| giriş satırı:", veri_satiri)
        if sayfalar[0] != "GİRİŞ BİLDİRİMLERİ":
            hatalar.append("ilk sayfa adı yanlış")
        if veri_satiri != 4:
            hatalar.append("excel giriş satır sayısı 4 değil")
        # John (satır 4) yabancı bilgileri dolu olmalı
        if ws.cell(row=4, column=8).value != "Alman":
            hatalar.append("John uyruk satırı Excel'de yok")
        if "tam bilgi zorunlu" in (ws.cell(row=4, column=13).value or ""):
            hatalar.append("John'un Excel notu gereksiz 'tam bilgi zorunlu' içeriyor")
        # Petra (satır 5) eksik bilgi uyarısı taşımalı
        if "DOĞUM TARİHİ" not in (ws.cell(row=5, column=13).value or ""):
            hatalar.append("Petra'nın Excel notunda eksik bilgi uyarısı yok")

    # 4) Gönderildi işaretle → bekleyenler azalmalı
    g1, _ = kbs.kbs_bekleyenler(DB, TAKIP)
    kesitler = [s["kesit"] for s in g1]
    kbs.kbs_markala(kesitler, "gonderildi", TAKIP)
    g2, c2 = kbs.kbs_bekleyenler(DB, TAKIP)
    print("Gönderildi işaretlendikten sonra — GİRİŞ:", len(g2), "| ÇIKIŞ:", len(c2))
    if len(g2) != 0:
        hatalar.append("işaretleme sonrası giriş bekleyen 0 değil")
    if len(c2) != 1:
        hatalar.append("çıkış bekleyen korunmadı (1 olmalı)")

    ozet = kbs.kbs_durum_ozet(DB, TAKIP)
    print("Özet:", ozet)
    if ozet["gonderilen"] != 4:
        hatalar.append("özet gonderilen 4 değil")

    print()
    if hatalar:
        print("HATALAR:")
        for h in hatalar:
            print(" -", h)
        raise SystemExit(1)
    print("TÜM TESTLER GEÇTİ ✔")
    print("Dosyalar:", DB, "|", TAKIP, "|", EXCEL)


if __name__ == "__main__":
    main()